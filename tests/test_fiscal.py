# -*- coding: utf-8 -*-
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_ROOT / "src"))

from fiscal import build_regras, validate_uf, write_regras  # noqa: E402


def _df(rows: list[dict[str, str]]) -> pd.DataFrame:
    base = {
        "NOME PRODUTO": "X", "NCM": "21069090", "CEST": "", "TRIBUTO": "TRIBUTADO",
        "IMPOSTO (% ICMS)": "0", "CFOP": "5101", "CST OU CSOSN": "102",
        "CST PIS": "49", "PIS CALCULO": "Percentual", "ALIQUOTA PIS": "0",
        "CST COFINS": "49", "COFINS CALCULO": "Percentual", "ALIQUOTA COFINS": "0",
        "CODIGO BENEFICIO FISCAL": "", "REDUCAO ICMS (%)": "0",
    }
    return pd.DataFrame([{**base, **r} for r in rows])


def test_agrupa_combinacoes_unicas() -> None:
    df = _df([{}, {}, {"NCM": "22030000", "CEST": "302100", "TRIBUTO": "ST", "CST OU CSOSN": "500"}])
    result = build_regras(df, "1")
    assert result.total_produtos == 3
    assert len(result.regras) == 2
    assert result.regras[1].cest == "0302100"
    assert result.regras[1].tributo == "S"


def test_simples_sem_pis_cofins_e_normal_com() -> None:
    assert build_regras(_df([{}]), "1").regras[0].pis is None
    assert build_regras(_df([{}]), "3").regras[0].pis == ("49", 0.0, None)


def test_rejeita_ncm_invalido() -> None:
    result = build_regras(_df([{"NCM": "123"}]), "1")
    assert not result.regras
    assert result.errors[0].field == "NCM"


def test_write_coluna_csosn_vs_cst(tmp_path: Path) -> None:
    template = _ROOT / "ImportacaoRegraNCMDadosFiscais.xlsx"
    regras = build_regras(_df([{}]), "1").regras

    out = tmp_path / "simples.xlsx"
    write_regras(regras, out, template, "1", "SP", "1")
    ws = openpyxl.load_workbook(out)["Dados Básicos"]
    assert (ws["G2"].value, ws["H2"].value) == (None, "102")

    out = tmp_path / "normal.xlsx"
    write_regras(regras, out, template, "3", "SP", "1")
    ws = openpyxl.load_workbook(out)["Dados Básicos"]
    assert (ws["G2"].value, ws["H2"].value) == ("102", None)


def test_uf_invalida() -> None:
    assert validate_uf(" sp ") == "SP"
    with pytest.raises(ValueError):
        validate_uf("XX")
